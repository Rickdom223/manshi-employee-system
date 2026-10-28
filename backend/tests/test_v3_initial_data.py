import importlib.util,json,os,sqlite3,subprocess,sys,tempfile
from pathlib import Path
from datetime import date,datetime
from decimal import Decimal
from contextlib import closing
from io import BytesIO
from openpyxl import Workbook,load_workbook
import pytest
from sqlalchemy import select,func
from sqlalchemy.orm import Session
from app.models import *
from app.core.database import Base,make_engine
from app.core.security import password_hash
from app.scripts.initialize_v3 import initialize
from app.services.initial_roster import DEFAULT_ROSTER,install_roster,read_roster,STATE_KEY
from test_v3_stage1 import login

ROOT=Path(__file__).resolve().parents[2]

def seed(db):initialize(db);return install_roster(db)

def test_all_document_fields_and_unique_admin(db):
    result=seed(db);_,records,_=read_roster()
    assert result['employees']==30 and result['accounts']==31 and result['system_admins']==1
    assert result['active_employees']==27 and result['departed_employees']==3
    assert result['active_monthly_base_salary']=='222900.00' and result['positions']==23 and result['supervisors']==4
    assert [(u.username,u.employee_id) for u in db.scalars(select(User).where(User.role=='SYS_ADMIN'))]==[('admin',None)]
    for expected in records:
        e=db.scalar(select(Employee).where(Employee.employee_no==expected.employee_no));u=db.scalar(select(User).where(User.employee_id==e.id));p=db.get(Position,e.position_id)
        assert e.full_name==u.username==expected.name
        assert (db.get(Department,e.department_id).name,p.name,p.grade)==(expected.department,expected.position,expected.grade)
        assert (e.phone,e.hire_date,e.resignation_date)==(expected.phone,expected.hire_date,expected.left_date)
        assert Decimal(e.base_salary_cents)/100==expected.base_salary
        assert (u.role,u.is_active,u.disabled_reason)==(expected.role,expected.is_active,expected.disabled_reason)
        assert password_hash.verify('Aa123',u.password_hash)
    for model in [Salary,PunchRecord,LeaveRequest,Attendance,Payroll,ProductionTask,LoginSession]:assert db.scalar(select(func.count()).select_from(model))==0
    assert db.scalar(select(func.count(EmploymentEvent.id)))==3

def test_seed_once_preserves_changes_and_new_data(db):
    seed(db);u=db.scalar(select(User).where(User.username=='admin'));new_hash=password_hash.hash('New1');u.password_hash=new_hash
    e=db.scalar(select(Employee).where(Employee.employee_no=='6'));e.phone='19999999999';db.commit()
    initialize(db);result=install_roster(db)
    assert result['already_installed'] is True and u.password_hash==new_hash and e.phone=='19999999999'
    assert db.scalar(select(func.count(User.id)))==31

def test_excel_multiple_grades_phone_salary_and_salary_edit(client,db):
    seed(db);login(client,'李名')
    wb=Workbook();ws=wb.active;ws.append(['工号','姓名','部门','职位','入职日期','职级','联系电话','税前月基本工资（元）'])
    ws.append(['31','新增同事','制琴生产部','制琴师','2026-10-01','L2','19900000031','8300.50'])
    ws.append(['32','职级缺失','制琴生产部','制琴师','2026-10-01',None,'19900000032','8000'])
    ws.append(['33','错误职级','制琴生产部','制琴师','2026-10-01','L9','19900000033','8000'])
    output=BytesIO();wb.save(output)
    result=client.post('/api/employees/import',files={'file':('roster.xlsx',output.getvalue())}).json()
    assert result['success_count']==1 and '多个职级' in result['data'][1]['reason']
    employee=client.get('/api/employees?employee_no=31').json()['data'][0]
    assert employee['grade']=='L2' and employee['phone']=='19900000031' and employee['base_salary']=='8300.50'
    updated=client.patch(f"/api/employees/{employee['id']}",json={'base_salary':'8400.99'}).json()['data']
    assert updated['base_salary']=='8400.99' and updated['account_username']=='新增同事'
    exported=load_workbook(BytesIO(client.get('/api/employees/export?employee_no=31').content),data_only=True)
    exported_rows=list(exported.active.values);assert exported_rows[0][5:8]==('职级','联系电话','税前月基本工资（元）')
    assert exported_rows[1][5:8]==('L2','19900000031','8400.99');exported.close()

def test_initial_roles_login_boundaries_and_salary_privacy(client,db,monkeypatch):
    import app.routers.v3_dashboard as dashboard_module
    monkeypatch.setattr(dashboard_module,'business_now',lambda:datetime(2026,10,28))
    seed(db);login(client,'admin','Admin123')
    users=client.get('/api/admin/users').json()['data']
    assert len(users)==31 and not any(u['username'] in {'demo_admin','ljm'} for u in users)
    assert '9500' not in client.get('/api/admin/users').text
    assert client.get('/api/employees').status_code==403 and client.get('/api/salaries').status_code==403
    login(client,'李名');assert client.get('/api/auth/me').json()['data']['role']=='HR'
    body=client.get('/api/employees?employee_no=1').json()['data'][0]
    assert body['base_salary']=='9500.00' and body['phone']=='19900000001'
    dashboard=client.get('/api/dashboard/hr').json()['data']
    assert dashboard['active_employees']==27 and dashboard['departed_this_month']==2
    assert [d['value'] for d in dashboard['departments']]==[9,5,7,6]
    assert dashboard['grades']==[{'name':'L1','value':7},{'name':'L2','value':13},{'name':'L3','value':3},{'name':'L4','value':4}]
    for name in ['汪维','杜硕','满园','高子象']:
        login(client,name);me=client.get('/api/auth/me').json()['data'];assert me['role']=='SUPERVISOR'
        ds=client.get('/api/departments').json()['data'];own=next(d for d in ds if d['supervisor_user_id']==me['id'])
        response=client.get(f"/api/departments/{own['id']}/employees")
        assert response.status_code==200 and all('base_salary' not in row and 'phone' not in row for row in response.json()['data'])
    login(client,'韩涛');assert client.get('/api/me/employee').json()['data']['base_salary']=='8700.00'
    assert client.get('/api/employees/1').status_code==403 and client.get('/api/salaries').status_code==403

@pytest.mark.parametrize('name',['许卓','唐悠','冯敏'])
def test_departed_initial_accounts_are_disabled(client,db,name):
    seed(db);assert client.post('/api/auth/login',json={'username':name,'password':'Aa123'}).status_code==401
    login(client,'admin','Admin123');row=next(u for u in client.get('/api/admin/departed-users').json()['data'] if u['username']==name)
    assert row['disabled_reason']=='EMPLOYEE_LEFT' and row['is_active'] is False
    assert client.post(f"/api/admin/users/{row['id']}/enable").status_code==422

def test_no_partial_roster_on_account_failure(db,monkeypatch):
    initialize(db)
    import app.services.initial_roster as module
    original=module.create_employee
    def fail_after_some_rows(session,body,role):
        if body.employee_no=='5':raise RuntimeError('simulated account creation failure')
        return original(session,body,role)
    monkeypatch.setattr(module,'create_employee',fail_after_some_rows)
    with pytest.raises(RuntimeError):install_roster(db)
    assert db.scalar(select(func.count(Employee.id)))==db.scalar(select(func.count(Position.id)))==0
    assert db.scalar(select(func.count(User.id)))==1 and db.get(InitialDataState,STATE_KEY) is None

def test_default_cli_installs_once_and_password_survives(tmp_path):
    path=tmp_path/'default.sqlite3';env=dict(os.environ,DATABASE_PATH=str(path),INITIAL_DATA_MODE='roster',PYTHONIOENCODING='utf-8')
    def run():
        result=subprocess.run([sys.executable,'-m','app.scripts.init_db'],cwd=ROOT/'backend',env=env,capture_output=True,text=True,encoding='utf-8')
        assert result.returncode==0,result.stdout+result.stderr
    run()
    with closing(sqlite3.connect(path)) as db, db:
        assert db.execute('SELECT count(*) FROM employees').fetchone()[0]==30
        new_hash=password_hash.hash('Different1');db.execute("UPDATE users SET password_hash=? WHERE username='admin'",(new_hash,))
    run()
    with closing(sqlite3.connect(path)) as db:
        assert db.execute("SELECT password_hash FROM users WHERE username='admin'").fetchone()[0]==new_hash
        assert db.execute('SELECT count(*) FROM users').fetchone()[0]==31

def test_rebuild_backs_up_dirty_database_and_copies_zero_old_rows():
    spec=importlib.util.spec_from_file_location('reset_initial_data',ROOT/'scripts/reset_initial_data.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    directory_root=(ROOT/'backend/data').resolve()
    with tempfile.TemporaryDirectory(dir=directory_root,prefix='test-approved-reset-') as directory:
        assert Path(directory).resolve().parent==directory_root
        target=Path(directory)/'live.sqlite3';engine=make_engine(target);Base.metadata.create_all(engine)
        with Session(engine) as db:
            initialize(db);db.add(User(username='demo_admin',role='SYS_ADMIN',password_hash=password_hash.hash('Old1')));db.commit()
        engine.dispose()
        result=module.rebuild(target)
        backup=Path(result['backup'])
        try:
            assert result['old_rows_copied']==0 and result['started_from_empty']
            with closing(sqlite3.connect(target)) as db:
                assert db.execute('SELECT count(*) FROM employees').fetchone()[0]==30
                assert db.execute("SELECT username FROM users WHERE role='SYS_ADMIN'").fetchall()==[('admin',)]
                assert db.execute("SELECT count(*) FROM users WHERE username='demo_admin'").fetchone()[0]==0
            with closing(sqlite3.connect(backup)) as db:assert db.execute("SELECT count(*) FROM users WHERE username='demo_admin'").fetchone()[0]==1
            bad=Path(directory)/'bad.json';data=json.loads(DEFAULT_ROSTER.read_text(encoding='utf-8'));data['employees'][0]['phone']='invalid';bad.write_text(json.dumps(data),encoding='utf-8')
            original=target.read_bytes()
            with pytest.raises(ValueError):module.rebuild(target,bad)
            assert target.read_bytes()==original
        finally:
            assert backup.resolve().parent==(ROOT/'backend/backups').resolve()
            backup.unlink()
