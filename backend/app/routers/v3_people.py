from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote
from fastapi import APIRouter, Depends, Query, UploadFile, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from app.core.database import get_db
from app.core.security import current_user, require_admin, require_hr, require_employee, password_hash
from app.core.errors import fail
from app.models import User, Employee, Position, Department, DepartmentSupervisor, AuditLog, EmploymentEvent
from app.schemas.v3 import *
from app.services.v3 import *
from app.routers.auth import valid_password

router=APIRouter(prefix='/api')

def get_or_404(db,model,id):
    obj=db.get(model,id)
    if obj is None:fail(404,'NOT_FOUND','记录不存在')
    return obj

def commit(db):
    try:db.commit()
    except IntegrityError:
        db.rollback();fail(409,'CONFLICT','编号、名称或关联记录冲突，请刷新后重试')

def validate_merged(schema,values):
    try:return schema.model_validate(values)
    except ValidationError:fail(422,'INVALID_FIELDS','必填字段不能为空或格式不正确')

@router.get('/bootstrap/status')
def status(db:Session=Depends(get_db)):
    return {'data':{'departments_initialized':db.scalar(select(func.count(Department.id)))==4,'admin_initialized':bool(db.scalar(select(User.id).where(User.role=='SYS_ADMIN'))),'has_active_hr':available_hr(db)}}

@router.post('/admin/bootstrap-first-hr')
def bootstrap(body:Bootstrap,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    if available_hr(db):fail(409,'HR_EXISTS','已有启用的 HR，请由 HR 维护员工')
    if not body.position_id:
        if not body.position_name or not db.get(Department,body.department_id):fail(422,'POSITION_REQUIRED','请选择有效部门并填写初始化职位')
        p=Position(code=f'INIT-{body.employee_no}'[:20],name=body.position_name,grade=body.grade,department_id=body.department_id)
        db.add(p);db.flush();body.position_id=p.id
    e,u=create_employee(db,body,'HR')
    audit(db,actor,'BOOTSTRAP_HR','users',u.id)
    commit(db)
    return {'data':employee_view(db,e),'message':f'HR 初始化成功，登录用户名：{u.username}；初始密码使用统一值'}

@router.post('/admin/recover-hr')
def recover(body:RecoverInput,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    if available_hr(db):fail(409,'HR_EXISTS','已有可用 HR')
    e=get_or_404(db,Employee,body.employee_id)
    u=db.scalar(select(User).where(User.employee_id==e.id))
    if e.status!='active' or not u:fail(422,'INVALID_EMPLOYEE','必须选择已有账号的在职员工')
    u.role='HR';u.is_active=True;u.disabled_reason=None;unassign(db,u.id);revoke(db,u.id)
    audit(db,actor,'RECOVER_HR','users',u.id);commit(db)
    return {'data':user_view(db,u)}

@router.get('/departments')
@router.get('/admin/departments/supervisors')
def departments(user:User=Depends(current_user),db:Session=Depends(get_db)):
    result=[]
    for d in db.scalars(select(Department).order_by(Department.id)):
        s=db.get(DepartmentSupervisor,d.id);u=db.get(User,s.user_id) if s else None
        result.append(dict(id=d.id,code=d.code,name=d.name,active_count=db.scalar(select(func.count(Employee.id)).where(Employee.department_id==d.id,Employee.status=='active')),position_count=db.scalar(select(func.count(Position.id)).where(Position.department_id==d.id)),supervisor_user_id=u.id if u else None,supervisor_name=u.display_name or u.username if u else None))
    return {'data':result,'total':len(result)}

@router.get('/departments/{id}/employees')
def department_employees(id:int,user:User=Depends(current_user),db:Session=Depends(get_db),page_no:int=Query(1,alias='page',ge=1),page_size:int=Query(50,ge=1,le=100)):
    assignment=db.get(DepartmentSupervisor,id)
    if user.role!='HR' and not(user.role=='SUPERVISOR' and assignment and assignment.user_id==user.id):fail(403,'FORBIDDEN','只允许查看所负责部门人员')
    return page(select(Employee).where(Employee.department_id==id,Employee.status=='active').order_by(Employee.id),db,page_no,page_size,lambda e:employee_view(db,e,True))

@router.put('/admin/departments/{id}/supervisor')
def supervisor(id:int,body:SupervisorInput,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    get_or_404(db,Department,id);u=get_or_404(db,User,body.user_id);e=db.get(Employee,u.employee_id) if u.employee_id else None
    if not u.is_active or not e or e.status!='active' or e.department_id!=id:fail(422,'INVALID_SUPERVISOR','主管必须为本部门启用账号的在职员工')
    old=db.get(DepartmentSupervisor,id)
    if old:db.delete(old);db.flush()
    unassign(db,u.id);u.role='SUPERVISOR';revoke(db,u.id)
    db.add(DepartmentSupervisor(department_id=id,user_id=u.id))
    audit(db,actor,'ASSIGN_SUPERVISOR','departments',id,f'用户 {u.id}');commit(db)
    return {'data':{'assigned':True}}

@router.delete('/admin/departments/{id}/supervisor')
def remove_supervisor(id:int,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    s=db.get(DepartmentSupervisor,id)
    if s:db.delete(s)
    audit(db,actor,'REMOVE_SUPERVISOR','departments',id);commit(db)
    return {'data':{'removed':True}}

def position_view(p):
    return dict(id=p.id,code=p.code,name=p.name,department_id=p.department_id,grade=p.grade,is_active=p.is_active,description=p.description)

@router.get('/positions')
def positions(department_id:int|None=None,grade:str|None=None,user:User=Depends(current_user),db:Session=Depends(get_db),page_no:int=Query(1,alias='page',ge=1),page_size:int=Query(100,ge=1,le=100)):
    q=select(Position).order_by(Position.id)
    if department_id:q=q.where(Position.department_id==department_id)
    if grade:q=q.where(Position.grade==grade)
    return page(q,db,page_no,page_size,position_view)

@router.get('/positions/{id}')
def position(id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
    return {'data':position_view(get_or_404(db,Position,id))}

@router.post('/positions')
def position_add(body:PositionInput,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    get_or_404(db,Department,body.department_id)
    p=Position(**body.model_dump());db.add(p);commit(db)
    return {'data':position_view(p)}

@router.patch('/positions/{id}')
def position_edit(id:int,body:PositionPatch,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    p=get_or_404(db,Position,id)
    body=validate_merged(PositionInput,{**{k:v for k,v in position_view(p).items() if k!='id'},**body.model_dump(exclude_unset=True)})
    get_or_404(db,Department,body.department_id)
    if p.department_id!=body.department_id and db.scalar(select(Employee.id).where(Employee.position_id==id)):fail(409,'POSITION_IN_USE','已被员工引用的职位不能更换所属部门，请新建职位')
    for k,v in body.model_dump().items():setattr(p,k,v)
    commit(db);return {'data':position_view(p)}

def employee_query(search=None,employee_no=None,department_id=None,position_id=None,status=None):
    q=select(Employee).order_by(Employee.id)
    if search:q=q.where(Employee.full_name.contains(search))
    if employee_no:q=q.where(Employee.employee_no==employee_no)
    if department_id:q=q.where(Employee.department_id==department_id)
    if position_id:q=q.where(Employee.position_id==position_id)
    if status:q=q.where(Employee.status=='active' if status=='ACTIVE' else Employee.status=='resigned')
    return q

@router.get('/employees')
def employees(search:str|None=None,employee_no:str|None=None,department_id:int|None=None,position_id:int|None=None,status:Literal['ACTIVE','LEFT']|None=None,page_no:int=Query(1,alias='page',ge=1),page_size:int=Query(50,ge=1,le=100),actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    return page(employee_query(search,employee_no,department_id,position_id,status),db,page_no,page_size,lambda e:employee_view(db,e))

def xlsx(wb,name):
    output=BytesIO();wb.save(output);output.seek(0)
    return StreamingResponse(output,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f"attachment; filename*=UTF-8''{quote(name)}"})

HEADERS=['工号','姓名','部门','职位','入职日期']
EXTRA_HEADERS=['职级','联系电话','税前月基本工资（元）']

@router.get('/employees/import-template')
def template(actor:User=Depends(require_hr)):
    wb=Workbook();ws=wb.active;ws.title='员工导入';ws.append(HEADERS+EXTRA_HEADERS)
    note=wb.create_sheet('填写说明');note.append(['每行填写真实员工；部门和职位必须已存在。入职日期使用 YYYY-MM-DD。工号已存在时拒绝覆盖。'])
    note.append(['示例（不参与导入）','MS002','张三','制琴生产部','制琴师','2026-10-01'])
    return xlsx(wb,'员工导入模板.xlsx')

@router.post('/employees/import')
async def import_employees(file:UploadFile,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith('.xlsx'):fail(422,'FILE_TYPE','仅支持 .xlsx 文件')
    data=await file.read(5*1024*1024+1)
    if not data or len(data)>5*1024*1024:fail(422,'FILE_SIZE','文件为空或超过5MB')
    import zipfile
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist())>25*1024*1024:raise ValueError()
        wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
        ws=wb.worksheets[0]
        if ws.max_row>1001:fail(422,'ROW_LIMIT','最多导入1000行')
        rows=iter(ws.values);header=list(next(rows))
        if header[:5]!=HEADERS:fail(422,'HEADERS','请使用本系统模板，前五列须为工号、姓名、部门、职位、入职日期')
    except HTTPException:raise
    except Exception:fail(422,'INVALID_EXCEL','无法读取文件，请使用有效的 Excel 模板')
    results=[]
    for line,row in enumerate(rows,2):
        if not any(v is not None for v in row):continue
        number=row[0] if row else None
        try:
            with db.begin_nested():
                d=db.scalar(select(Department).where(Department.name==str(row[2]).strip()))
                extras={str(key).strip():row[index] for index,key in enumerate(header) if index<len(row)}
                grade=str(extras.get('职级') or '').strip()
                q=select(Position).where(Position.name==str(row[3]).strip(),Position.department_id==d.id,Position.is_active==True) if d else None
                if q is not None and grade:q=q.where(Position.grade==grade)
                matches=db.scalars(q).all() if q is not None else []
                if len(matches)>1:raise ValueError('同名职位有多个职级，请填写职级列')
                p=matches[0] if matches else None
                if not d or not p:raise ValueError('部门不存在或职位/职级不属于该部门')
                day=row[4].date() if hasattr(row[4],'hour') else row[4]
                salary=str(extras.get('税前月基本工资（元）') or '0').replace(',','').strip()
                phone=extras.get('联系电话')
                e,u=create_employee(db,EmployeeInput(employee_no=str(number).strip(),name=row[1],department_id=d.id,position_id=p.id,hire_date=day,phone=str(phone).strip() if phone else None,base_salary=salary))
                audit(db,actor,'IMPORT_EMPLOYEE','employees',e.id)
                results.append(dict(row=line,employee_no=number,success=True,username=u.username,reason='初始密码使用统一值'))
        except Exception as exc:
            message=exc.detail['message'] if isinstance(exc,HTTPException) else ('工号或用户名冲突' if isinstance(exc,IntegrityError) else str(exc))
            results.append(dict(row=line,employee_no=number,success=False,reason=message))
    wb.close();commit(db)
    if not results:fail(422,'EMPTY_FILE','文件没有员工数据')
    return {'data':results,'success_count':sum(r['success'] for r in results)}

@router.get('/employees/export')
def export(search:str|None=None,employee_no:str|None=None,department_id:int|None=None,position_id:int|None=None,status:Literal['ACTIVE','LEFT']|None=None,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    wb=Workbook();ws=wb.active;ws.append(HEADERS+EXTRA_HEADERS+['在职状态','离职日期'])
    for e in db.scalars(employee_query(search,employee_no,department_id,position_id,status)):
        v=employee_view(db,e);ws.append([e.employee_no,e.full_name,v['department_name'],v['position_name'],str(e.hire_date),v['grade'],e.phone,v['base_salary'],v['status'],str(e.resignation_date or '')])
    return xlsx(wb,'员工档案.xlsx')

@router.post('/employees')
def employee_add(body:EmployeeInput,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    e,u=create_employee(db,body);audit(db,actor,'CREATE_EMPLOYEE','employees',e.id);commit(db)
    return {'data':employee_view(db,e),'message':f'建档成功，登录用户名：{u.username}；初始密码使用统一值'}

@router.get('/me/employee')
def my_employee(user:User=Depends(require_employee),db:Session=Depends(get_db)):
    return {'data':employee_view(db,db.get(Employee,user.employee_id))}

@router.get('/employees/{id}')
def employee(id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
    e=get_or_404(db,Employee,id);s=db.get(DepartmentSupervisor,e.department_id)
    own=user.employee_id==e.id
    if user.role!='HR' and not own and not(user.role=='SUPERVISOR' and s and s.user_id==user.id):fail(403,'FORBIDDEN','无权查看该员工')
    return {'data':employee_view(db,e,user.role!='HR' and not own)}

@router.patch('/employees/{id}')
def employee_edit(id:int,body:EmployeePatch,actor:User=Depends(require_hr),db:Session=Depends(get_db)):
    e=get_or_404(db,Employee,id)
    current=dict(employee_no=e.employee_no,name=e.full_name,department_id=e.department_id,position_id=e.position_id,hire_date=e.hire_date,phone=e.phone,remarks=e.remark,base_salary=Decimal(e.base_salary_cents)/100)
    body=validate_merged(EmployeeInput,{**current,**body.model_dump(exclude_unset=True)})
    p=get_or_404(db,Position,body.position_id)
    if p.department_id!=body.department_id or (p.id!=e.position_id and not p.is_active):fail(422,'INVALID_POSITION','职位与部门不匹配或已停用')
    u=db.scalar(select(User).where(User.employee_id==e.id))
    if u and e.department_id!=body.department_id:unassign(db,u.id);revoke(db,u.id)
    if e.status=='resigned' and e.resignation_date<body.hire_date:fail(422,'INVALID_DATE','入职日期晚于离职日期')
    for k,v in body.model_dump().items():setattr(e,{'name':'full_name','remarks':'remark','base_salary':'base_salary_cents'}.get(k,k),int(v*100) if k=='base_salary' else v)
    if u:u.display_name=body.name
    audit(db,actor,'UPDATE_EMPLOYEE','employees',id);commit(db)
    return {'data':employee_view(db,e)}

def users_query(db,search,department_id,role,is_active,departed):
    q=select(User).outerjoin(Employee,Employee.id==User.employee_id).order_by(User.id)
    if search:q=q.where(or_(User.username.contains(search),Employee.full_name.contains(search),Employee.employee_no.contains(search)))
    if department_id:q=q.where(Employee.department_id==department_id)
    if role:q=q.where(User.role==role)
    if is_active is not None:q=q.where(User.is_active==is_active)
    if departed:q=q.where(Employee.status=='resigned')
    return q

@router.get('/admin/users')
@router.get('/admin/departed-users')
def users(request:__import__('fastapi').Request,search:str|None=None,department_id:int|None=None,role:Literal['SYS_ADMIN','HR','SUPERVISOR','EMPLOYEE']|None=None,is_active:bool|None=None,left_from:date|None=None,left_to:date|None=None,page_no:int=Query(1,alias='page',ge=1),page_size:int=Query(50,ge=1,le=100),actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    q=users_query(db,search,department_id,role,is_active,request.url.path.endswith('departed-users'))
    if left_from:q=q.where(Employee.resignation_date>=left_from)
    if left_to:q=q.where(Employee.resignation_date<=left_to)
    return page(q,db,page_no,page_size,lambda u:user_view(db,u))

@router.get('/admin/users/{id}')
def user_detail(id:int,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=get_or_404(db,User,id)
    events=db.scalars(select(EmploymentEvent).where(EmploymentEvent.employee_id==u.employee_id).order_by(EmploymentEvent.id)).all() if u.employee_id else []
    return {'data':dict(user_view(db,u),employment_events=[dict(id=e.id,event_type=e.event_type,event_date=e.event_date,note=e.note) for e in events])}

def mutable_user(db,id):
    u=get_or_404(db,User,id)
    if u.role=='SYS_ADMIN':fail(403,'ADMIN_PROTECTED','系统管理员账号受保护，请通过本人账号设置修改密码')
    return u

@router.patch('/admin/users/{id}/role')
def role(id:int,body:RoleInput,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=mutable_user(db,id);e=db.get(Employee,u.employee_id)
    if not e or e.status!='active' or not u.is_active:fail(422,'INACTIVE','只允许为启用的在职员工授权')
    old=u.role;u.role=body.role
    if old!=body.role:unassign(db,u.id);revoke(db,u.id)
    audit(db,actor,'CHANGE_ROLE','users',id,f'{old} → {body.role}');commit(db)
    return {'data':user_view(db,u)}

@router.patch('/admin/users/{id}/username')
def username(id:int,body:UsernameInput,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=mutable_user(db,id);u.username=body.username;revoke(db,id)
    audit(db,actor,'CHANGE_USERNAME','users',id);commit(db);return {'data':user_view(db,u)}

@router.post('/admin/users/{id}/disable')
def disable(id:int,body:ReasonInput,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=mutable_user(db,id);u.is_active=False;u.disabled_reason=body.reason;revoke(db,id);unassign(db,id)
    audit(db,actor,'DISABLE_USER','users',id,body.reason);commit(db);return {'data':user_view(db,u)}

@router.post('/admin/users/{id}/enable')
def enable(id:int,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=mutable_user(db,id);e=db.get(Employee,u.employee_id)
    if not e or e.status!='active':fail(422,'EMPLOYEE_LEFT','员工仍为离职状态，请先由 HR 办理复职')
    u.is_active=True;u.disabled_reason=None
    audit(db,actor,'ENABLE_USER','users',id);commit(db);return {'data':user_view(db,u)}

@router.post('/admin/users/{id}/reset-password')
def reset(id:int,body:ResetPassword,actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    u=mutable_user(db,id);valid_password(body.new_password);u.password_hash=password_hash.hash(body.new_password);revoke(db,id)
    audit(db,actor,'RESET_PASSWORD','users',id);commit(db);return {'data':{'changed':True}}

@router.get('/admin/audit-logs')
def logs(action:str|None=None,page_no:int=Query(1,alias='page',ge=1),page_size:int=Query(50,ge=1,le=100),actor:User=Depends(require_admin),db:Session=Depends(get_db)):
    q=select(AuditLog).order_by(AuditLog.id.desc())
    if action:q=q.where(AuditLog.action==action)
    return page(q,db,page_no,page_size,lambda r:dict(id=r.id,actor_user_id=r.actor_user_id,action=r.action,target_type=r.target_type,target_id=r.target_id,summary=r.summary,created_at=r.created_at))
