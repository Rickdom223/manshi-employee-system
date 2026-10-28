from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select,func
from app.models import TaskMember,Inspection,ProductionTask,Employee
from test_auth import admin,auth
from test_employees import dictionaries
from test_spec_stage2 import staff,staff_client

def test_TASK_01_02_03_reports(client,auth,staff,staff_client,db):
    original=db.get(Employee,staff['id'])
    other=Employee(employee_no='TASK_OTHER',full_name='质检测试员工乙',department_id=original.department_id,position_id=original.position_id,hire_date=original.hire_date,phone=None)
    db.add(other);db.commit()
    body={'name':'MS-01民谣吉他生产任务','guitar_model':'MS-01','planned_quantity':5,'completed_quantity':0,'owner_employee_id':staff['id'],'participant_ids':[staff['id'],other.id],'start_date':'2026-10-27','due_date':'2026-11-15','priority':'medium'}
    response=client.post('/api/tasks',headers=auth,json=body);assert response.status_code==201,response.text
    task=response.json()['data'];id=task['id']
    assert len(task['participant_ids'])==2
    assert staff_client.get('/api/employee/tasks').json()['total']==1
    assert staff_client.get('/api/tasks/'+str(id)).status_code==403
    assert staff_client.post(f'/api/tasks/{id}/transition',json={'status':'in_progress'}).status_code==403
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'in_progress'}).status_code==200
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'pending_qc'}).status_code==422
    body['completed_quantity']=5
    assert client.patch(f'/api/tasks/{id}',headers=auth,json=body).status_code==200
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'pending_qc'}).status_code==200
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'completed'}).status_code==409
    assert client.post(f'/api/tasks/{id}/inspections',headers=auth,json={'result':'failed','notes':'返工'}).json()['data']['status']=='in_progress'
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'pending_qc'}).status_code==200
    assert client.post(f'/api/tasks/{id}/inspections',headers=auth,json={'result':'passed','notes':'音准检查通过'}).status_code==201
    assert client.post(f'/api/tasks/{id}/transition',headers=auth,json={'status':'completed'}).status_code==200
    assert client.patch(f'/api/tasks/{id}',headers=auth,json=body).status_code==409
    assert len(client.get(f'/api/tasks/{id}/inspections',headers=auth).json()['data'])==2
    assert client.get('/api/dashboard',headers=auth).json()['data']['in_progress_tasks']==0
    for type in ['departments','attendance','payroll','production']:
        response=client.get(f'/api/reports/{type}?month=2026-10',headers=auth)
        assert response.status_code==200,response.text
        export=client.get(f'/api/reports/export?type={type}&month=2026-10',headers=auth)
        assert export.status_code==200
        assert load_workbook(BytesIO(export.content)).active.max_row>=1
        assert staff_client.get(f'/api/reports/{type}?month=2026-10').status_code==403

def test_task_nonparticipant_and_invalid(client,auth,staff,staff_client,db):
    original=db.get(Employee,staff['id'])
    other=Employee(employee_no='NO_MEMBER',full_name='不相关员工',department_id=original.department_id,position_id=original.position_id,hire_date=original.hire_date)
    db.add(other);db.commit()
    body={'name':'非本人任务','guitar_model':'MODEL','planned_quantity':1,'owner_employee_id':other.id,'participant_ids':[],'start_date':'2026-10-27','due_date':'2026-11-01'}
    assert client.post('/api/tasks',headers=auth,json=body).status_code==201
    assert staff_client.get('/api/employee/tasks?employee_id='+str(other.id)).json()['total']==0
    assert client.post('/api/tasks',headers=auth,json=body|{'participant_ids':[staff['id'],staff['id']]}).status_code==422
    other.status='resigned';other.resignation_date=other.hire_date;db.commit()
    assert client.post('/api/tasks',headers=auth,json=body).status_code==422

def test_dashboard_statistics(client,auth,staff):
    data=client.get('/api/dashboard',headers=auth).json()['data']
    assert data['active']==1 and len(data['attendance_trend'])==7
    assert all(r['count']==0 for r in data['attendance_trend'])
    assert sum(r['count'] for r in data['department_distribution'])==1
