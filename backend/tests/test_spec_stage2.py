from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select,func
from openpyxl import load_workbook
from app.main import app
from app.models import Employee,User,Attendance,LeaveRequest,Payroll
from test_auth import admin,auth
from test_employees import dictionaries

@pytest.fixture
def staff(client,auth,dictionaries):
    d,_,p,_,_=dictionaries
    response=client.post('/api/employees',headers=auth,json={'name':'测试员工甲','phone':'13800000001','hire_date':'2026-10-01','department_id':d,'position_id':p,'base_salary':'6000.00','create_account':True,'initial_password':'Test1234'})
    assert response.status_code==201,response.text
    return response.json()['data']

@pytest.fixture
def staff_client(staff):
    with TestClient(app,headers={'Origin':'http://testserver'}) as c:
        response=c.post('/api/auth/login',json={'username':staff['employee_code'],'password':'Test1234'})
        assert response.status_code==200,response.text
        c.headers['X-CSRF-Token']=response.json()['data']['csrf_token']
        yield c

def test_EMP_01_permissions_departure(client,auth,staff,staff_client,db):
    assert len(staff['employee_code'])<=10
    assert client.post('/api/employees',headers=auth,json={'name':'缺基本工资','phone':'13800000001','department_id':staff['department_id'],'position_id':staff['position_id'],'hire_date':'2026-10-01'}).status_code==422
    assert staff_client.get('/api/auth/me').json()['data']['role']=='EMPLOYEE'
    for path in ['/api/employees','/api/departments','/api/dashboard','/api/payroll','/api/leaves']:
        assert staff_client.get(path).status_code==403
    assert staff_client.patch('/api/employee/me',json={'phone':'13900000001','employee_id':999}).status_code==422
    assert staff_client.patch('/api/employee/me',json={'phone':'13900000001'}).status_code==200
    assert staff_client.get('/api/employee/me?employee_id=999').json()['data']['id']==staff['id']
    assert client.post(f"/api/employees/{staff['id']}/enable-account",headers=auth,json={'initial_password':'Test1234'}).status_code==409
    assert client.post(f"/api/employees/{staff['id']}/leave-company",headers=auth,json={'departure_date':'2026-10-27'}).status_code==200
    assert staff_client.get('/api/employee/me').status_code==401
    assert db.get(Employee,staff['id'])

def test_LEAVE_01_02_03(client,auth,staff,staff_client,db):
    request={'leave_type':'personal','start_date':'2026-10-30','end_date':'2026-11-02','reason':'测试请假'}
    result=staff_client.post('/api/employee/leaves',json=request)
    assert result.status_code==201,result.text
    id=result.json()['data']['id']
    assert client.post(f'/api/leaves/{id}/approve',headers=auth).status_code==200
    records=db.scalars(select(Attendance).where(Attendance.employee_id==staff['id'])).all()
    assert len(records)==2 and all(r.status=='leave' and r.source=='leave' for r in records)
    assert client.post(f'/api/leaves/{id}/approve',headers=auth).status_code==409
    assert staff_client.post(f'/api/employee/leaves/{id}/cancel').status_code==409
    duplicate=staff_client.post('/api/employee/leaves',json=request).json()['data']['id']
    assert client.post(f'/api/leaves/{duplicate}/approve',headers=auth).status_code==409
    assert db.get(LeaveRequest,duplicate).status=='pending'
    assert db.scalar(select(func.count(Attendance.id)))==2
    assert client.post('/api/attendance',headers=auth,json={'employee_id':staff['id'],'attendance_date':'2026-10-29','status':'normal'}).status_code==201
    manual_conflict=staff_client.post('/api/employee/leaves',json=request|{'start_date':'2026-10-28','end_date':'2026-10-29'}).json()['data']['id']
    assert client.post(f'/api/leaves/{manual_conflict}/approve',headers=auth).status_code==409
    assert db.get(LeaveRequest,manual_conflict).status=='pending'
    assert db.scalar(select(func.count(Attendance.id)))==3
    assert client.patch(f'/api/attendance/{records[0].id}',headers=auth,json={'employee_id':staff['id'],'attendance_date':'2026-10-30','status':'normal'}).status_code==409
    other=LeaveRequest(employee_id=staff['id']+100,leave_type='sick',start_date=records[0].attendance_date,end_date=records[0].attendance_date,reason='不存在员工')
    # A different real employee's leave cannot be cancelled by this account.
    original=db.get(Employee,staff['id'])
    employee=Employee(employee_no='OTHER',full_name='测试员工乙',phone=None,department_id=original.department_id,position_id=original.position_id,hire_date=original.hire_date)
    db.add(employee);db.flush();other.employee_id=employee.id;db.add(other);db.commit()
    assert staff_client.post(f'/api/employee/leaves/{other.id}/cancel').status_code==403
    assert staff_client.post('/api/employee/leaves',json=request|{'start_date':'2026-10-31','end_date':'2026-11-01'}).status_code==422

def test_PAY_01_02_03(client,auth,staff,staff_client,db):
    assert client.post('/api/payroll/generate',headers=auth,json={'salary_month':'2026-10'}).json()['data']['created']==1
    row=client.get('/api/payroll?month=2026-10',headers=auth).json()['data'][0];id=row['id']
    response=client.patch(f'/api/payroll/{id}',headers=auth,json={'bonus':'123.45','deduction':'23.44'})
    assert response.json()['data']['net_salary']=='6100.01'
    assert client.post('/api/payroll/generate',headers=auth,json={'salary_month':'2026-10'}).json()['data']['created']==0
    assert db.get(Payroll,id).bonus_cents==12345
    assert client.patch(f'/api/payroll/{id}',headers=auth,json={'bonus':'0','deduction':'99999'}).status_code==422
    assert client.post(f'/api/payroll/{id}/mark-paid',headers=auth).status_code==200
    assert client.patch(f'/api/payroll/{id}',headers=auth,json={'bonus':'1','deduction':'0'}).status_code==409
    assert client.post(f'/api/payroll/{id}/reopen',headers=auth).status_code==200
    assert staff_client.get('/api/payroll').status_code==403
    own=staff_client.get('/api/employee/payroll?employee_id=999').json()['data']
    assert all(r['employee_id']==staff['id'] for r in own)
    client.patch(f"/api/employees/{staff['id']}",headers=auth,json={'base_salary':'8000'})
    assert db.get(Payroll,id).base_salary_cents==600000

def test_FILE_01(client,auth,staff,staff_client):
    client.post('/api/payroll/generate',headers=auth,json={'salary_month':'2026-10'})
    for url in ['/api/employees/export','/api/attendance/export?month=2026-10','/api/payroll/export?month=2026-10']:
        result=client.get(url,headers=auth);assert result.status_code==200,result.text
        workbook=load_workbook(BytesIO(result.content));assert workbook.active.cell(1,1).value=='员工编号'
        assert staff_client.get(url).status_code==403
